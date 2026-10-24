"""读取与Java相同的受控夹具；仅为候选比较，不用于平台入库。"""

import json
import sys
import time
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader, __version__
from pypdf.errors import FileNotDecryptedError, PdfReadError, PdfStreamError

directory = Path(sys.argv[1])
started = time.monotonic()
result = {"pypdf": __version__, "pages": [], "rejected": []}
pdf = PdfReader(directory / "native.pdf", strict=True)
assert len(pdf.pages) == 3
for number, page in enumerate(pdf.pages, start=1):
    text = page.extract_text()
    (directory / f"python-page-{number}.txt").write_text(text)
    if number == 2:
        assert not text.strip()
    else:
        assert ("Alpha policy" if number == 1 else "Rotated policy") in text
    result["pages"].append(
        {"page": number, "rotation": page.rotation, "empty": not text.strip()}
    )
for name in ("encrypted.pdf", "malformed.pdf"):
    try:
        invalid = PdfReader(directory / name, strict=True)
        invalid.pages[0].extract_text()
    except (FileNotDecryptedError, PdfReadError, PdfStreamError):
        result["rejected"].append(name)
    else:
        raise AssertionError(f"{name} must fail")

document = Document(directory / "structure.docx")
elements = list(document.iter_inner_content())
assert len(elements) == 5
assert document.paragraphs[0].text == "资料处理说明"
assert document.paragraphs[0].style.style_id == "Heading1"
assert document.paragraphs[1].text == ""
table = document.tables[0]
assert table.cell(0, 0).text == "项目"
assert table.rows[1].cells[0]._tc.tcPr.gridSpan.val == 2
body = []
for element in elements:
    if isinstance(element, Paragraph):
        body.append("P:" + element.text)
    elif isinstance(element, Table):
        for row in element.rows:
            # 公共row.cells展开合并格；探针读取物理XML格才能与Java原序一致。
            cells = row._tr.findall(qn("w:tc"))
            # 按w:t精确读取实际文本节点，避免自定义节点itertext重复文本。
            texts = [
                "\n".join(
                    "".join(t.text or "" for t in p.iter(qn("w:t")))
                    for p in cell.findall(qn("w:p"))
                )
                for cell in cells
            ]
            body.append("ROW:" + "|".join(texts))
(directory / "python-docx.txt").write_text("\n".join(body) + "\n")
assert (directory / "java-docx.txt").read_text() == (
    directory / "python-docx.txt"
).read_text()
for page in (1, 2, 3):
    # 不把各库的末尾换行差别误称同字节输出；正文可读性只比较trim后的内容。
    assert (directory / f"java-page-{page}.txt").read_text().strip() == (
        directory / f"python-page-{page}.txt"
    ).read_text().strip()
result["docx"] = "same ordered paragraphs and physical merged cells as Java"
result["parseMillis"] = round((time.monotonic() - started) * 1000, 2)
(directory / "python-result.json").write_text(
    json.dumps(result, ensure_ascii=False, indent=2) + "\n"
)
print(json.dumps(result, ensure_ascii=False))
